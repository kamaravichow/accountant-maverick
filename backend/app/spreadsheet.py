"""Server-side spreadsheet engine for the agent: read, write, add live formulas, and evaluate them.

The agent builds workpapers *with real Excel formulas* (so the accountant can audit and tweak them
in the browser spreadsheet editor or Excel), then evaluates them here with the ``formulas`` engine
to read back the computed numbers.
"""

from __future__ import annotations

import io
import os
import re
import tempfile
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries

INR_FORMAT = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D7DE")
MONEY_HINT = re.compile(r"(amount|value|tax|igst|cgst|sgst|cess|total|tds|debit|credit|balance|itc|difference|interest|"
                        r"fee|payable|refund|income|salary|cost|price|depreciation|wdv)", re.I)


def load(data: bytes, data_only: bool = False):
    return load_workbook(io.BytesIO(data), data_only=data_only)


def save(wb) -> bytes:
    try:
        wb.calculation.fullCalcOnLoad = True  # Excel/LibreOffice recalc formulas on open
    except Exception:
        pass
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _jsonable(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10] if isinstance(v, datetime) and v.time() == datetime.min.time() else v.isoformat()
    return v


def describe(data: bytes) -> dict:
    wb = load(data)
    return {"sheets": [{"name": ws.title, "dimensions": ws.dimensions, "max_row": ws.max_row, "max_col": ws.max_column}
                       for ws in wb.worksheets]}


def read_range(data: bytes, sheet: str | None = None, cell_range: str | None = None, max_rows: int = 200,
               evaluate: bool = True) -> dict:
    """Return values (and formulas) for a range. Formula results come from the evaluation engine,
    falling back to cached values stored in the file."""
    wb_f = load(data)
    ws_f = wb_f[sheet] if sheet else wb_f.worksheets[0]
    wb_v = load(data, data_only=True)
    ws_v = wb_v[ws_f.title]
    if cell_range:
        min_col, min_row, max_col, max_row_ = range_boundaries(cell_range)
    else:
        min_col, min_row, max_col, max_row_ = 1, 1, ws_f.max_column, ws_f.max_row
    max_row_ = min(max_row_, min_row + max_rows - 1)
    computed = evaluate_workbook(data) if evaluate and _has_formulas(ws_f, min_row, max_row_, min_col, max_col) else {}
    rows, formulas = [], {}
    for r in range(min_row, max_row_ + 1):
        row = []
        for c in range(min_col, max_col + 1):
            cell = ws_f.cell(r, c)
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                ref = f"{get_column_letter(c)}{r}"
                formulas[ref] = v
                v = computed.get((ws_f.title.upper(), ref), ws_v.cell(r, c).value)
            row.append(_jsonable(v))
        rows.append(row)
    return {"sheet": ws_f.title, "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row_}",
            "rows": rows, "formulas": formulas, "truncated": ws_f.max_row > max_row_ and not cell_range}


def _has_formulas(ws, r1, r2, c1, c2) -> bool:
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True):
        if any(isinstance(v, str) and v.startswith("=") for v in row):
            return True
    return False


def evaluate_workbook(data: bytes) -> dict[tuple[str, str], Any]:
    """Evaluate every formula with the `formulas` engine -> {(SHEET_UPPER, 'A1'): value}."""
    import formulas

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        model = formulas.ExcelModel().loads(path).finish()
        sol = model.calculate()
    finally:
        os.unlink(path)
    out: dict[tuple[str, str], Any] = {}
    for key, val in sol.items():
        m = re.match(r"'\[[^\]]+\](.+)'!([A-Z]+\d+)$", str(key))
        if not m:
            continue
        v = getattr(val, "value", val)
        try:
            v = v[0][0]
        except Exception:
            pass
        if hasattr(v, "item"):
            v = v.item()
        if type(v).__name__ in ("XlError",) or str(v).startswith("#"):
            v = str(v)
        out[(m.group(1).upper(), m.group(2))] = v
    return out


def write_table(data: bytes | None, sheet: str, headers: list[str], rows: list[list], start_cell: str = "A1",
                title: str | None = None, totals: list[str] | None = None, replace_sheet: bool = True) -> bytes:
    """Write a formatted table. ``totals`` = header names to total with live =SUBTOTAL formulas."""
    wb = load(data) if data else Workbook()
    if not data:
        wb.remove(wb.active)
    if sheet in wb.sheetnames and replace_sheet:
        del wb[sheet]
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    col0, row0, _, _ = range_boundaries(f"{start_cell}:{start_cell}")
    r = row0
    if title:
        ws.cell(r, col0, title).font = Font(bold=True, size=13)
        r += 2
    for j, h in enumerate(headers):
        c = ws.cell(r, col0 + j, h)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, Alignment(wrap_text=True, vertical="center")
        c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    header_row = r
    for row in rows:
        r += 1
        for j, v in enumerate(row):
            c = ws.cell(r, col0 + j, v)
            if isinstance(v, (int, float)) and not isinstance(v, bool) and MONEY_HINT.search(str(headers[j] if j < len(headers) else "")):
                c.number_format = INR_FORMAT
    first, last = header_row + 1, r
    if totals and rows:
        r += 1
        ws.cell(r, col0, "Total").font = Font(bold=True)
        for name in totals:
            if name in headers:
                j = headers.index(name)
                col = get_column_letter(col0 + j)
                c = ws.cell(r, col0 + j, f"=SUBTOTAL(9,{col}{first}:{col}{last})")
                c.font, c.number_format = Font(bold=True), INR_FORMAT
    for j, h in enumerate(headers):
        width = max([len(str(h))] + [len(str(rw[j])) for rw in rows[:200] if j < len(rw) and rw[j] is not None])
        ws.column_dimensions[get_column_letter(col0 + j)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = ws.cell(header_row + 1, col0)
    if rows:
        ws.auto_filter.ref = f"{get_column_letter(col0)}{header_row}:{get_column_letter(col0 + len(headers) - 1)}{last}"
    return save(wb)


def set_cells(data: bytes | None, sheet: str, cells: dict[str, Any], number_format: str | None = None) -> bytes:
    """Set values or formulas (strings starting with '=') on individual cells."""
    wb = load(data) if data else Workbook()
    if not data:
        wb.active.title = sheet
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    for ref, v in cells.items():
        c = ws[ref]
        c.value = v
        if number_format == "inr" or (number_format is None and isinstance(v, (int, float)) and abs(v) >= 1000):
            c.number_format = INR_FORMAT
        elif number_format:
            c.number_format = number_format
    return save(wb)


def append_rows(data: bytes | None, sheet: str, headers: list[str], rows: list[dict]) -> bytes:
    """Append dict rows to a register sheet, creating it with headers if needed (column order preserved)."""
    if data:
        wb = load(data)
    else:
        wb = Workbook()
        wb.active.title = sheet
    if sheet not in wb.sheetnames:
        wb.create_sheet(sheet)
    ws = wb[sheet]
    existing = [c.value for c in ws[1]] if ws.max_row >= 1 and ws.cell(1, 1).value is not None else []
    if not existing:
        for j, h in enumerate(headers, 1):
            c = ws.cell(1, j, h)
            c.fill, c.font = HEADER_FILL, HEADER_FONT
        existing = headers
        ws.freeze_panes = "A2"
    for h in headers:
        if h not in existing:
            existing.append(h)
            ws.cell(1, len(existing), h).font = HEADER_FONT
            ws.cell(1, len(existing)).fill = HEADER_FILL
    for rec in rows:
        r = ws.max_row + 1
        for j, h in enumerate(existing, 1):
            v = rec.get(h)
            c = ws.cell(r, j, v)
            if isinstance(v, (int, float)) and not isinstance(v, bool) and MONEY_HINT.search(h):
                c.number_format = INR_FORMAT
    return save(wb)


def records_from_sheet(data: bytes, sheet: str | None = None) -> list[dict]:
    wb = load(data, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else f"col_{i + 1}" for i, h in enumerate(rows[0])]
    return [{h: _jsonable(v) for h, v in zip(headers, r)} for r in rows[1:] if any(v is not None for v in r)]
