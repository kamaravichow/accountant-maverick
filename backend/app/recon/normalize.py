"""Data cleaning for messy accounting exports (Tally, bank statements, GST portal Excel, client spreadsheets).

Handles: Indian digit grouping (1,23,456.78), Dr/Cr suffixes, (brackets) negatives, '-' / 'nil' blanks,
₹/Rs/INR prefixes, many date formats incl. Excel serials, header synonyms, junk title rows above the
real header, repeated page headers inside PDFs, and duplicate rows.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Iterable

_AMOUNT_JUNK = re.compile(r"(₹|rs\.?|inr|\s)", re.I)


def parse_amount(value) -> Decimal | None:
    """'1,23,456.50 Dr' -> -123456.50 ; '(500)' -> -500 ; '5,000 Cr' -> 5000 ; '-' -> None."""
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(value))
    s = str(value).strip()
    if s in ("", "-", "--", "nil", "NIL", "Nil", "NA", "N/A", "null", "None"):
        return None
    neg = False
    low = s.lower()
    if low.endswith("dr") or low.endswith("dr."):
        neg = True
        s = re.sub(r"dr\.?$", "", s, flags=re.I)
    elif low.endswith("cr") or low.endswith("cr."):
        s = re.sub(r"cr\.?$", "", s, flags=re.I)
    s = _AMOUNT_JUNK.sub("", s)
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s.endswith("-"):  # trailing minus from some bank exports
        neg, s = True, s[:-1]
    s = s.replace(",", "")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    return -d if neg else d


_DATE_FORMATS = [
    "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y", "%d.%m.%y", "%Y-%m-%d", "%Y/%m/%d",
    "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d %b %y", "%d-%B-%Y", "%d %B %Y", "%b %d, %Y", "%d%m%Y",
    "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
]


def parse_date(value, dayfirst: bool = True) -> date | None:
    """Indian formats are day-first. Excel serial numbers are supported."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < float(value) < 80000:
        return date(1899, 12, 30) + timedelta(days=int(value))
    s = str(value).strip()
    if re.fullmatch(r"\d{5}(\.0+)?", s):
        return date(1899, 12, 30) + timedelta(days=int(float(s)))
    s = re.sub(r"\s+", " ", s)
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# Canonical column -> synonyms seen in the wild (lowercased, punctuation stripped)
HEADER_SYNONYMS = {
    "date": ["date", "txn date", "transaction date", "value date", "posting date", "tran date", "voucher date",
             "invoice date", "inv date", "bill date", "document date", "doc date"],
    "narration": ["narration", "description", "particulars", "remarks", "transaction details", "details",
                  "transaction remarks", "narrative"],
    "reference": ["chq no", "cheque no", "chq ref no", "ref no", "reference", "reference no", "utr", "chq/ref no",
                  "cheque/reference no", "ref no./cheque no.", "instrument no", "chqno"],
    "debit": ["debit", "withdrawal", "withdrawals", "withdrawal amt", "withdrawal amount", "dr", "debit amount",
              "debits", "paid out", "amount dr"],
    "credit": ["credit", "deposit", "deposits", "deposit amt", "deposit amount", "cr", "credit amount", "credits",
               "paid in", "amount cr"],
    "amount": ["amount", "txn amount", "transaction amount", "amt"],
    "balance": ["balance", "closing balance", "running balance", "available balance", "bal"],
    "gstin": ["gstin", "gstin of supplier", "gstin/uin", "supplier gstin", "gstin of recipient", "party gstin",
              "gstin/uin of supplier", "ctin", "gst no", "gst number"],
    "party": ["party", "party name", "supplier", "supplier name", "vendor", "vendor name", "trade/legal name",
              "trade name", "legal name", "customer", "customer name", "name", "ledger", "ledger name"],
    "invoice_no": ["invoice no", "invoice number", "inv no", "bill no", "bill number", "voucher no", "document number",
                   "doc no", "invoice details invoice number", "supplier invoice no", "vch no"],
    "taxable_value": ["taxable value", "taxable amount", "assessable value", "taxable val", "txval", "basic amount"],
    "igst": ["igst", "integrated tax", "igst amount", "integrated tax amount"],
    "cgst": ["cgst", "central tax", "cgst amount", "central tax amount"],
    "sgst": ["sgst", "sgst/utgst", "state/ut tax", "state tax", "sgst amount", "utgst"],
    "cess": ["cess", "cess amount"],
    "invoice_value": ["invoice value", "total", "total amount", "gross amount", "bill amount", "invoice amount", "val"],
    "rate": ["rate", "gst rate", "tax rate", "rate %", "rate (%)", "rt"],
    "hsn": ["hsn", "hsn code", "hsn/sac", "sac", "hsn sac", "hsn/sac code"],
    "place_of_supply": ["place of supply", "pos"],
    "tan": ["tan", "tan of deductor"],
    "tds": ["tds", "tax deducted", "tds amount", "tax deducted/collected", "tds deposited"],
    "section": ["section", "tds section", "nature of payment"],
    "quantity": ["qty", "quantity"],
}


def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9/ ]", "", str(h or "").lower().replace("\n", " ").replace("_", " "))).strip()


def map_headers(headers: list) -> dict[int, str]:
    """Map column index -> canonical name; unknown columns keep a cleaned version of their header."""
    out: dict[int, str] = {}
    used: set[str] = set()
    normed = [_norm_header(h) for h in headers]
    # exact synonym matches first, then 'contains' matches
    for pass_ in ("exact", "contains"):
        for i, h in enumerate(normed):
            if i in out or not h:
                continue
            for canon, syns in HEADER_SYNONYMS.items():
                if canon in used:
                    continue
                if (pass_ == "exact" and h in syns) or (
                    pass_ == "contains" and any(len(s) > 3 and s in h for s in syns)
                ):
                    out[i] = canon
                    used.add(canon)
                    break
    for i, h in enumerate(normed):
        if i not in out:
            out[i] = re.sub(r"\s+", "_", h) or f"col_{i + 1}"
    return out


def find_header_row(rows: list[list], max_scan: int = 30) -> int:
    """Statements often have logos, addresses and account info above the real header row."""
    best, best_score = 0, -1
    for i, row in enumerate(rows[:max_scan]):
        mapped = map_headers(row)
        score = sum(1 for v in mapped.values() if v in HEADER_SYNONYMS)
        if score > best_score:
            best, best_score = i, score
    return best


def rows_to_records(rows: list[list], header_row: int | None = None) -> list[dict]:
    if not rows:
        return []
    hr = find_header_row(rows) if header_row is None else header_row
    header = rows[hr]
    mapping = map_headers(header)
    header_norm = [_norm_header(h) for h in header]
    records = []
    for row in rows[hr + 1:]:
        if not any(str(c).strip() for c in row if c is not None):
            continue
        if [_norm_header(c) for c in row[: len(header_norm)]] == header_norm:
            continue  # repeated header on each PDF page
        rec = {mapping.get(i, f"col_{i + 1}"): (v.strip() if isinstance(v, str) else v) for i, v in enumerate(row)}
        records.append(rec)
    return records


NUMERIC_FIELDS = {"debit", "credit", "amount", "balance", "taxable_value", "igst", "cgst", "sgst", "cess",
                  "invoice_value", "rate", "tds", "quantity"}
DATE_FIELDS = {"date"}


def clean_records(records: Iterable[dict]) -> tuple[list[dict], dict]:
    """Type-convert known fields, derive signed ``amount`` for statements, drop totals/opening lines, dedupe."""
    cleaned, seen = [], set()
    stats = {"input": 0, "dropped_blank_or_total": 0, "duplicates": 0, "unparsed_dates": 0}
    for rec in records:
        stats["input"] += 1
        r = dict(rec)
        text = " ".join(str(v) for v in r.values() if v is not None).lower()
        if re.match(r"\s*(opening balance|closing balance|grand total|total\b|sub total|b/f|c/f|carried forward|brought forward)",
                     str(r.get("narration") or r.get("party") or "").lower()) or text.strip() == "":
            stats["dropped_blank_or_total"] += 1
            continue
        for k in list(r):
            if k in NUMERIC_FIELDS:
                v = parse_amount(r[k])
                r[k] = float(v) if v is not None else None
            elif k in DATE_FIELDS:
                d = parse_date(r[k])
                if d is None and r[k] not in (None, ""):
                    stats["unparsed_dates"] += 1
                r[k] = d.isoformat() if d else None
            elif k == "gstin" and r[k]:
                r[k] = str(r[k]).strip().upper().replace(" ", "")
        if "amount" not in r and ("debit" in r or "credit" in r):
            r["amount"] = (r.get("credit") or 0) - (r.get("debit") or 0)
        if r.get("date") is None and "date" in r and not any(r.get(k) for k in ("debit", "credit", "amount")):
            # wrapped narration continuation line in bank PDFs - append to previous
            if cleaned and r.get("narration"):
                cleaned[-1]["narration"] = f"{cleaned[-1].get('narration') or ''} {r['narration']}".strip()
            stats["dropped_blank_or_total"] += 1
            continue
        key = tuple(sorted((k, str(v)) for k, v in r.items()))
        if key in seen:
            stats["duplicates"] += 1
            continue
        seen.add(key)
        cleaned.append(r)
    stats["output"] = len(cleaned)
    return cleaned, stats


_FY_TOKEN = re.compile(r"(?<!\d)(?:20)?(\d{2})\s*[-/_]\s*(?:20)?(\d{2})(?!\d)")


def _strip_fy(m: re.Match) -> str:
    a, b = int(m.group(1)), int(m.group(2))
    return "" if b == (a + 1) % 100 else m.group(0)  # only consecutive years = FY token (25-26, 2025/26)


def normalize_invoice_no(inv) -> str:
    """'INV/2025-26/000123' and 'INV/25-26/123' -> 'INV123' (drop FY tokens, separators, leading zeros)."""
    s = str(inv or "").upper().strip()
    s = _FY_TOKEN.sub(_strip_fy, s)
    s = re.sub(r"\bFY\b", "", s)
    s = re.sub(r"[^A-Z0-9]", "", s)
    s = re.sub(r"(?<![0-9])0+(?=[0-9])", "", s)
    return s


# ------------------------------------------------------------------ file readers

def read_table_bytes(data: bytes, filename: str, sheet: str | None = None) -> list[list]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|") if text.strip() else csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    if name.endswith(".pdf"):
        return pdf_tables(data)
    raise ValueError(f"Unsupported table format: {filename} (use .xlsx, .csv or .pdf; convert .xls to .xlsx)")


def pdf_tables(data: bytes) -> list[list]:
    import pdfplumber

    rows: list[list] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            if not tables:
                tables = page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"}) or []
            for t in tables:
                rows.extend([[(c or "").replace("\n", " ").strip() for c in r] for r in t])
    return rows


def load_records(data: bytes, filename: str, sheet: str | None = None) -> tuple[list[dict], dict]:
    rows = read_table_bytes(data, filename, sheet)
    return clean_records(rows_to_records(rows))
